import os
from typing import Optional, Dict, List, Any
from docx import Document as DocxDocument
from docx.shared import Cm, Pt


class MyDocument:
    def __init__(self, filename: Optional[str] = None):
        self.doc = DocxDocument(filename) if filename else DocxDocument()
        self.heading_counters: Dict[int, int] = {}
        self._section = self.doc.sections[0]

    def apply_visual_settings(
        self,
        top_margin: float = 2.0,
        bottom_margin: float = 2.0,
        left_margin: float = 2.0,
        right_margin: float = 1.0,
        font_name: str = 'Calibri',
        font_size: int = 11
    ) -> None:
        self._section.top_margin = Cm(top_margin)
        self._section.bottom_margin = Cm(bottom_margin)
        self._section.left_margin = Cm(left_margin)
        self._section.right_margin = Cm(right_margin)

        style = self.doc.styles['Normal']
        style.font.name = font_name
        style.font.size = Pt(font_size)

        for level in range(1, 4):
            try:
                heading_style = self.doc.styles[f'Heading {level}']
                heading_style.font.name = font_name
            except KeyError:
                pass

    def add_heading_numbered(self, text: str, level: int = 1) -> Any:
        self.heading_counters[level] = self.heading_counters.get(level, 0) + 1
        
        for lvl in list(self.heading_counters.keys()):
            if lvl > level:
                self.heading_counters[lvl] = 0

        numbers = [str(self.heading_counters.get(i, 0)) for i in range(1, level + 1)]
        full_heading = f'{".".join(numbers)}. {text}'
        return self.doc.add_heading(full_heading, level=level)

    def add_paragraph(self, text: str, style: Optional[str] = None) -> Any:
        return self.doc.add_paragraph(text, style=style) if style else self.doc.add_paragraph(text)

    def add_page_break(self) -> None:
        self.doc.add_page_break()

    def add_table(
        self, 
        data: List[List[Any]], 
        headers: List[str], 
        title: Optional[str] = None
    ) -> None:
        if title:
            self.add_paragraph(title)

        if not data:
            self.add_paragraph("Veri yok")
            return

        table = self.doc.add_table(rows=len(data) + 1, cols=len(headers))
        table.style = 'Table Grid'

        hdr_cells = table.rows[0].cells
        for i, header in enumerate(headers):
            run = hdr_cells[i].paragraphs[0].add_run(str(header))
            run.bold = True

        for row_idx, row_data in enumerate(data):
            row_cells = table.rows[row_idx + 1].cells
            for col_idx, val in enumerate(row_data):
                row_cells[col_idx].text = str(val)

    def add_image(self, image_source: Any, width_cm: float = 15.0) -> None:
        """image_source: Dosya yolu (str) veya BytesIO nesnesi olabilir."""
        if hasattr(image_source, "seek"):
            image_source.seek(0)
            self.doc.add_picture(image_source, width=Cm(width_cm))
        elif isinstance(image_source, str) and os.path.exists(image_source):
            self.doc.add_picture(image_source, width=Cm(width_cm))
        else:
            self.add_paragraph("Resim yüklenemedi.")

    def save_report(self, path: str) -> bool:
        try:
            dir_name = os.path.dirname(os.path.abspath(path))
            if dir_name:
                os.makedirs(dir_name, exist_ok=True)
            self.doc.save(path)
            return True
        except Exception:
            return False

    def add_section_break(self) -> None:
        self._section = self.doc.add_section()

    def clear(self) -> None:
        for p in list(self.doc.paragraphs):
            p._element.getparent().remove(p._element)
        for t in list(self.doc.tables):
            t._element.getparent().remove(t._element)
        self.heading_counters.clear()