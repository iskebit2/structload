#utils/report_dataframe.py
from __future__ import annotations

from io import BytesIO
from typing import Any, Optional

import pandas as pd


class ReportDataFrame(pd.DataFrame):
    _metadata = [
        "custom_title",
        "custom_desc",
        "column_descriptions",
        "column_units",
        "column_formats",
        "_table_style",
        "graphics",
    ]

    def __init__(self, *args, **kwargs):
        self.custom_title = kwargs.pop("custom_title", None)
        self.custom_desc = kwargs.pop("custom_desc", None)
        self.column_descriptions = kwargs.pop("column_descriptions", {})
        self.column_units = kwargs.pop("column_units", {})
        self.column_formats = kwargs.pop("column_formats", {})
        self._table_style = kwargs.pop("table_style", "Table Grid")
        self.graphics = kwargs.pop("graphics", [])

        super().__init__(*args, **kwargs)

    @property
    def _constructor(self):
        def _c(*args, **kwargs):
            df = ReportDataFrame(*args, **kwargs)
            df.custom_title = self.custom_title
            df.custom_desc = self.custom_desc
            df.column_descriptions = self.column_descriptions.copy()
            df.column_units = self.column_units.copy()
            df.column_formats = self.column_formats.copy()
            df._table_style = self._table_style
            df.graphics = [graphic.copy() for graphic in self.graphics]
            return df

        return _c

    def set_column_description(
        self, column: str, description: str
    ) -> ReportDataFrame:
        self.column_descriptions[column] = description
        return self

    def set_column_unit(self, column: str, unit: str) -> ReportDataFrame:
        self.column_units[column] = unit
        return self

    def set_column_format(self, column: str, fmt: str) -> ReportDataFrame:
        self.column_formats[column] = fmt
        return self

    def add_graphic(
        self,
        data: BytesIO | bytes,
        title: Optional[str] = None,
        description: Optional[str] = None,
        width: Optional[float] = None,
        height: Optional[float] = None,
        format: str = "png",
    ) -> ReportDataFrame:
        if isinstance(data, bytes):
            data = BytesIO(data)

        if not isinstance(data, BytesIO):
            raise TypeError("Graphic data BytesIO veya bytes olmalıdır.")

        self.graphics.append(
            {
                "data": data,
                "title": title,
                "description": description,
                "width": width,
                "height": height,
                "format": format,
            }
        )
        return self

    def clear_graphics(self) -> ReportDataFrame:
        self.graphics.clear()
        return self

    def _format_value(self, column: str, value: Any) -> str:
        if pd.isna(value):
            return ""

        fmt = self.column_formats.get(column)
        if fmt is None:
            return str(value)

        try:
            return format(value, fmt)
        except (TypeError, ValueError):
            return str(value)

    def _format_row(self, row: pd.Series) -> list[str]:
        return [
            self._format_value(column, val) for column, val in row.items()
        ]

    def round_numeric(self, decimals: int = 2) -> ReportDataFrame:
        numeric_cols = self.select_dtypes(include="number").columns
        self[numeric_cols] = self[numeric_cols].round(decimals)
        return self

    def to_markdown(self, *args, **kwargs) -> str:
        lines = []

        if self.custom_title:
            lines.append(f"## {self.custom_title}")

        if self.custom_desc:
            lines.append(f"*{self.custom_desc}*")

        if self.empty:
            lines.append("Gösterilecek veri bulunamadı.")
            return "\n".join(lines)

        cols = [str(col) for col in self.columns]
        lines.append("| " + " | ".join(cols) + " |")
        lines.append("| " + " | ".join("---" for _ in cols) + " |")

        if self.column_descriptions:
            descs = [self.column_descriptions.get(col, "") for col in self.columns]
            lines.append("| " + " | ".join(descs) + " |")

        if self.column_units:
            units = [self.column_units.get(col, "") for col in self.columns]
            lines.append("| " + " | ".join(units) + " |")

        for _, row in self.iterrows():
            lines.append("| " + " | ".join(self._format_row(row)) + " |")

        return "\n".join(lines)

    def _repr_html_(self) -> str:
        from html import escape

        parts = []

        if self.custom_title:
            parts.append(f"<h4>{escape(str(self.custom_title))}</h4>")

        if self.custom_desc:
            parts.append(f"<p>{escape(str(self.custom_desc))}</p>")

        if self.empty:
            parts.append("<p>Gösterilecek veri bulunamadı.</p>")
            return "".join(parts)

        html = ["<table><thead><tr>"]
        for col in self.columns:
            html.append(f"<th>{escape(str(col))}</th>")
        html.append("</tr>")

        if self.column_descriptions:
            html.append("<tr>")
            for col in self.columns:
                desc = self.column_descriptions.get(col, "")
                html.append(f"<th>{escape(str(desc))}</th>")
            html.append("</tr>")

        if self.column_units:
            html.append("<tr>")
            for col in self.columns:
                unit = self.column_units.get(col, "")
                html.append(f"<th>{escape(str(unit))}</th>")
            html.append("</tr>")

        html.append("</thead><tbody>")

        for _, row in self.iterrows():
            html.append("<tr>")
            for col, val in row.items():
                text = self._format_value(col, val)
                html.append(f"<td>{escape(text)}</td>")
            html.append("</tr>")

        html.append("tbody></table>")
        parts.append("".join(html))

        return "".join(parts)

    def save_to_docx(
        self,
        document,
        title: Optional[str] = None,
        desc: Optional[str] = None,
        level: int = 2,
    ) -> None:
        from docx.shared import Inches, Pt

        final_title = title if title is not None else self.custom_title
        final_desc = desc if desc is not None else self.custom_desc

        if final_title:
            document.add_heading(final_title, level=level)

        if final_desc:
            p = document.add_paragraph(final_desc)
            p.style = document.styles["Normal"]

        if self.empty:
            document.add_paragraph("Gösterilecek veri bulunamadı.")
            return

        has_descriptions = bool(self.column_descriptions)
        has_units = bool(self.column_units)
        extra_rows = int(has_descriptions) + int(has_units)

        table = document.add_table(
            rows=1 + extra_rows + len(self),
            cols=len(self.columns),
        )
        table.style = self._table_style

        row_index = 0

        for col_index, col in enumerate(self.columns):
            cell = table.rows[row_index].cells[col_index]
            cell.text = str(col)
            for run in cell.paragraphs[0].runs:
                run.bold = True
                run.font.size = Pt(9)
        row_index += 1

        if has_descriptions:
            for col_index, col in enumerate(self.columns):
                table.rows[row_index].cells[col_index].text = str(
                    self.column_descriptions.get(col, "")
                )
            row_index += 1

        if has_units:
            for col_index, col in enumerate(self.columns):
                table.rows[row_index].cells[col_index].text = str(
                    self.column_units.get(col, "")
                )
            row_index += 1

        for _, data_row in self.iterrows():
            values = self._format_row(data_row)
            for col_index, value in enumerate(values):
                table.rows[row_index].cells[col_index].text = value
            row_index += 1

        for graphic in self.graphics:
            graphic_title = graphic.get("title")
            graphic_desc = graphic.get("description")

            if graphic_title:
                p = document.add_paragraph()
                p.add_run(graphic_title).bold = True

            if graphic_desc:
                document.add_paragraph(graphic_desc)

            data = graphic["data"]
            data.seek(0)

            width = Inches(graphic["width"]) if graphic.get("width") else None
            height = Inches(graphic["height"]) if graphic.get("height") else None

            document.add_picture(data, width=width, height=height)
            document.add_paragraph()

        document.add_paragraph()

    def print(self) -> None:
        if self.custom_title:
            print(f"\n{self.custom_title}")
            print("-" * len(self.custom_title))

        if self.custom_desc:
            print(f"{self.custom_desc}\n")

        if self.empty:
            print("Veri yok")
            return

        print(" | ".join(str(col) for col in self.columns))

        if self.column_descriptions:
            print(
                " | ".join(
                    self.column_descriptions.get(col, "") for col in self.columns
                )
            )

        if self.column_units:
            print(
                " | ".join(
                    self.column_units.get(col, "") for col in self.columns
                )
            )

        for _, row in self.iterrows():
            print(" | ".join(self._format_row(row)))